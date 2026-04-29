#!/usr/bin/env python3
"""
GÖKHİSAR Rapor - 4. YÖNTEM bölümü DOCX oluşturucu
"""

from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_shading(cell, color):
    """Hücre arka plan rengini ayarla"""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shading)

def add_heading(doc, text, level):
    """Başlık ekle"""
    heading = doc.add_heading(text, level=level)
    return heading

def add_table(doc, headers, rows):
    """Tablo ekle"""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    
    # Header
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        header_cells[i].text = header
        set_cell_shading(header_cells[i], "CCCCCC")
        header_cells[i].paragraphs[0].runs[0].bold = True
    
    # Rows
    for row_data in rows:
        row_cells = table.add_row().cells
        for i, cell_data in enumerate(row_data):
            row_cells[i].text = str(cell_data)
    
    return table

def create_report():
    doc = Document()
    
    # Stil ayarları
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    
    # ==================== 4. YÖNTEM ====================
    add_heading(doc, "4. YÖNTEM", 1)
    
    # 4.1 Sistem Genel Mimarisi
    add_heading(doc, "4.1 Sistem Genel Mimarisi", 2)
    
    doc.add_paragraph(
        "GÖKHİSAR Yer Kontrol İstasyonu, PySide6 (Qt6) tabanlı bir masaüstü uygulaması "
        "olarak geliştirilmiştir. Sistem, çok katmanlı (multi-tier) bir mimari üzerine "
        "inşa edilmiş olup aşağıdaki ana bileşenlerden oluşmaktadır:"
    )
    
    add_table(doc, 
        ["Katman", "Bileşen", "Sorumluluk"],
        [
            ["UI Layer", "MainWindow, StatusPanel, ControlPanel, VideoDisplay, LogPanel", "Kullanıcı arayüzü ve etkileşim"],
            ["Worker Layer", "UDPVideoWorker, TCPCommandWorker", "Arka plan iş parçacıkları (thread)"],
            ["Communication Layer", "UDP/TCP Socket", "Ağ haberleşmesi"],
            ["Hardware Layer", "Raspberry Pi + Servo + Kamera", "Fiziksel kontrol"]
        ]
    )
    
    doc.add_paragraph()
    
    # 4.1.1 Thread Yapısı
    add_heading(doc, "4.1.1 Thread Yapısı", 3)
    
    doc.add_paragraph(
        "Sistemde asenkron işlem için QThread tabanlı Worker sınıfları kullanılmaktadır:"
    )
    
    bullet_points = [
        "Ana Thread (UI Thread): Kullanıcı etkileşimlerini yönetir",
        "UDP Worker Thread: Video akışını alır (non-blocking)",
        "TCP Worker Thread: Komut iletişimini yönetir (güvenilir iletim)"
    ]
    for point in bullet_points:
        p = doc.add_paragraph(point, style='List Bullet')
    
    doc.add_paragraph(
        "Signal/Slot mekanizması ile thread'ler arası güvenli (thread-safe) iletişim sağlanmaktadır."
    )
    
    # ==================== AŞAMA 1 ====================
    add_heading(doc, "4.2 Aşama 1: Hedef Tespit Süreci", 2)
    
    add_heading(doc, "4.2.1 Activity Diagram Açıklaması", 3)
    
    doc.add_paragraph(
        "Hedef tespit süreci aşağıdaki adımlardan oluşmaktadır:\n\n"
        "1. UDP Video Akışı Başlatılır\n"
        "2. Frame alınır ve görüntü işleme yapılır\n"
        "3. Hedef tespit edilirse koordinatlar hesaplanır\n"
        "4. IFF (Identification Friend or Foe) kontrolü yapılır\n"
        "5. Status Panel güncellenir ve log kaydedilir\n"
        "6. Aşama 2'ye (Takip) geçilir"
    )
    
    add_heading(doc, "4.2.2 Sequence Diagram Açıklaması", 3)
    
    doc.add_paragraph(
        "Hedef tespit sürecinde bileşenler arası iletişim:\n\n"
        "1. Kamera (RPI) → UDPVideoWorker: UDP Frame gönderilir\n"
        "2. UDPVideoWorker → VideoDisplay: frame_received Signal'ı tetiklenir\n"
        "3. VideoDisplay → StatusPanel: Hedef bilgisi güncellenir\n"
        "4. StatusPanel → LogPanel: Olay kaydedilir"
    )
    
    add_heading(doc, "4.2.3 Teknik Detaylar", 3)
    
    add_table(doc,
        ["Parametre", "Değer", "Açıklama"],
        [
            ["Protokol", "UDP", "Düşük gecikme, yüksek hız"],
            ["Buffer Boyutu", "65535 byte", "Maksimum UDP datagram"],
            ["Video Format", "JPEG", "Sıkıştırılmış frame"],
            ["Port", "5000", "Varsayılan video portu"]
        ]
    )
    
    doc.add_paragraph()
    
    # ==================== AŞAMA 2 ====================
    add_heading(doc, "4.3 Aşama 2: Hedef Takip Süreci", 2)
    
    add_heading(doc, "4.3.1 Activity Diagram Açıklaması", 3)
    
    doc.add_paragraph(
        "Hedef takip süreci:\n\n"
        "1. Hedef tespit edildikten sonra menzil kontrolü yapılır\n"
        "2. Menzil bandına göre (YAKIN/ORTA/UZAK) aksiyon belirlenir\n"
        "3. Servo pozisyonu hesaplanır ve komut gönderilir\n"
        "4. Takip kilidi sağlanır\n"
        "5. Hedef angajman menziline girdiğinde Aşama 3'e geçilir"
    )
    
    add_heading(doc, "4.3.2 Menzil Bantları", 3)
    
    add_table(doc,
        ["Bant", "Mesafe", "Durum", "Aksiyon"],
        [
            ["YAKIN", "0-5m", "Kritik", "Angajman öncelikli"],
            ["ORTA", "5-10m", "Takip", "Aktif izleme"],
            ["UZAK", "10-15m", "İzleme", "Pasif gözlem"],
            ["DIŞI", ">15m", "Dışı", "Kayıt dışı"]
        ]
    )
    
    doc.add_paragraph()
    
    add_heading(doc, "4.3.3 Servo Kontrol Algoritması", 3)
    
    doc.add_paragraph(
        "Servo komutları 8-byte binary paket formatında gönderilmektedir:"
    )
    
    add_table(doc,
        ["Alan", "Boyut", "Tip", "Açıklama"],
        [
            ["CMD_ID", "1 byte", "uint8", "Komut tanımlayıcı (1=SERVO)"],
            ["X", "2 byte", "int16", "Azimut açısı (-180° ile +180°)"],
            ["Y", "2 byte", "int16", "Elevasyon açısı (-90° ile +90°)"],
            ["FLAGS", "1 byte", "uint8", "Ek bayraklar"],
            ["CHECKSUM", "2 byte", "uint16", "Tüm byte'ların toplamı mod 65536"]
        ]
    )
    
    doc.add_paragraph()
    
    # ==================== AŞAMA 3 ====================
    add_heading(doc, "4.4 Aşama 3: Angajman (Ateş Etme) Süreci", 2)
    
    add_heading(doc, "4.4.1 Activity Diagram Açıklaması", 3)
    
    doc.add_paragraph(
        "Angajman süreci:\n\n"
        "1. Hedef takip altına alındıktan sonra çalışma modu kontrol edilir\n"
        "2. Moda göre karar mekanizması işletilir:\n"
        "   - MANUEL: Operatör onayı beklenir\n"
        "   - YARI_OTONOM: Sistem önerisi + Operatör onayı\n"
        "   - TAM_OTONOM: Otomatik karar (IFF kontrolü ile)\n"
        "3. Güvenlik kilidi kontrolü yapılır\n"
        "4. IFF son kontrolü yapılır (DOST ise ateş engellenir)\n"
        "5. Ateş komutu TCP üzerinden gönderilir\n"
        "6. ACK beklenir ve sonuç loglanır"
    )
    
    add_heading(doc, "4.4.2 Güvenlik Mekanizmaları", 3)
    
    doc.add_paragraph("Sistemde çok katmanlı güvenlik uygulanmaktadır:")
    
    safety_points = [
        "Fiziksel Kilit: Ateş butonu varsayılan olarak devre dışı",
        "Yazılımsal Kilit: _fire_unlocked flag kontrolü",
        "IFF Kontrolü: Dost/Düşman tanımlama",
        "Acil Durdur: Tüm işlemleri anında durdurur",
        "Otomatik Kilit: Ateş sonrası kilit otomatik kapanır"
    ]
    for i, point in enumerate(safety_points, 1):
        doc.add_paragraph(f"{i}. {point}")
    
    add_heading(doc, "4.4.3 Ateş Komutu Protokolü", 3)
    
    doc.add_paragraph(
        "Ateş paketi CMD_ID=2 ve FLAGS=1 değerleri ile gönderilir. "
        "X ve Y değerleri 0 olarak ayarlanır (mevcut pozisyon kullanılır)."
    )
    
    # ==================== STATE MACHINE ====================
    add_heading(doc, "4.5 State Machine Diagram: Çalışma Modları", 2)
    
    doc.add_paragraph(
        "Sistem açıldığında varsayılan olarak MANUEL modda başlar. "
        "Üç farklı çalışma modu bulunmaktadır:"
    )
    
    add_table(doc,
        ["Mod", "Servo Kontrolü", "Ateş Kontrolü", "Takip"],
        [
            ["MANUEL", "Operatör kontrolünde", "Manuel onay gerekli", "Manuel"],
            ["YARI_OTONOM", "Sistem yardımlı", "Operatör onayı gerekli", "Otomatik"],
            ["TAM_OTONOM", "Otomatik", "Otomatik (IFF ile)", "Otomatik"]
        ]
    )
    
    doc.add_paragraph()
    
    doc.add_paragraph(
        "ACİL DURDUR (EMERGENCY_STOP) durumunda tüm komutlar engellenir, "
        "servo devre dışı bırakılır ve ateş sistemi kilitlenir."
    )
    
    # ==================== HABERLEŞMe ====================
    add_heading(doc, "4.6 Haberleşme Mimarisi", 2)
    
    add_heading(doc, "4.6.1 Network Topolojisi", 3)
    
    doc.add_paragraph(
        "Sistem iki ana haberleşme kanalı kullanmaktadır:\n\n"
        "1. UDP (Port 5000): Kameradan video akışı (Raspberry Pi → Yer Kontrol İstasyonu)\n"
        "2. TCP (Port 5001): Komut iletişimi (Yer Kontrol İstasyonu → Raspberry Pi)"
    )
    
    add_heading(doc, "4.6.2 Protokol Karşılaştırması", 3)
    
    add_table(doc,
        ["Özellik", "UDP (Video)", "TCP (Komut)"],
        [
            ["Güvenilirlik", "Best-effort", "Garantili"],
            ["Sıralama", "Garanti yok", "Sıralı"],
            ["Gecikme", "Düşük", "Orta"],
            ["Kullanım", "Video stream", "Kritik komutlar"],
            ["Handshake", "Yok", "3-way handshake"]
        ]
    )
    
    doc.add_paragraph()
    
    # ==================== SONUÇ ====================
    add_heading(doc, "4.7 Sonuç", 2)
    
    doc.add_paragraph(
        "GÖKHİSAR sistemi, modüler mimari, thread-safe iletişim ve çok katmanlı güvenlik "
        "prensipleri üzerine inşa edilmiştir. Üç farklı çalışma modu (Manuel, Yarı Otonom, "
        "Tam Otonom) ile esnek operasyon imkanı sağlanmaktadır. UDP protokolü ile düşük "
        "gecikmeli video akışı, TCP protokolü ile güvenilir komut iletimi gerçekleştirilmektedir."
    )
    
    # Kaydet
    output_path = "/home/kagan/Gokhisar/docs/4_YONTEM_Raporu.docx"
    doc.save(output_path)
    print(f"✅ Rapor oluşturuldu: {output_path}")
    return output_path

if __name__ == "__main__":
    create_report()
